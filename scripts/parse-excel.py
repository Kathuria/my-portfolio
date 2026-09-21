#!/usr/bin/env python3
"""
Excel Parser for Flight Data

Reads flights.xlsx and outputs flights-temp.json for the TypeScript generator.
This is a helper script that extracts raw data without processing.
"""

import openpyxl
import json
import sys
from pathlib import Path

def parse_excel():
    xlsx_path = Path('flights.xlsx')
    
    if not xlsx_path.exists():
        print(f"❌ Error: {xlsx_path} not found")
        sys.exit(1)
    
    # Load workbook
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    
    # Parse header
    header = [cell.value for cell in ws[1]]
    print(f"📋 Found columns: {header}")
    
    # Parse data rows
    flights = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:  # Skip empty rows
            flights.append({
                'origin': row[0],
                'destination': row[1],
                'layover': row[2] if len(row) > 2 else None,
                'airline': row[3] if len(row) > 3 else 'Unknown',
                'distanceKm': float(row[4]) if len(row) > 4 and row[4] else 0,
                'distanceMiles': float(row[5]) if len(row) > 5 and row[5] else 0,
            })
    
    # Write temp JSON
    output_path = Path('flights-temp.json')
    with open(output_path, 'w') as f:
        json.dump(flights, f, indent=2)
    
    print(f"✅ Parsed {len(flights)} flights")
    print(f"💾 Written to: {output_path}")

if __name__ == '__main__':
    parse_excel()
